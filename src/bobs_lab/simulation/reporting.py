"""导出平均排名汇总 Excel 与运行元数据."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def save_summary(
    columns: dict[str, dict[int, float | None]],
    path: Path,
    metadata: dict[str, str | int | None],
) -> None:
    """导出模式/MMR 平均排名矩阵及 metadata 工作表."""
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary = workbook.active
    if summary is None:
        raise RuntimeError("无法创建 Excel 汇总工作表.")
    summary.title = "average_position"
    summary.append(["候选数量", *columns])
    counts = sorted({count for values in columns.values() for count in values})
    for count in counts:
        summary.append([count, *(values.get(count) for values in columns.values())])
    for row in summary.iter_rows(min_row=2, min_col=2):
        for cell in row:
            cell.number_format = "0.0000"

    info = workbook.create_sheet("metadata")
    info.append(["key", "value"])
    for key, value in metadata.items():
        info.append([key, "random" if key == "seed" and value is None else value])
    for sheet in workbook:
        sheet.freeze_panes = "B2" if sheet == summary else "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for index, cells in enumerate(sheet.columns, start=1):
            width = max(len(str(cell.value or "")) for cell in cells) + 2
            sheet.column_dimensions[get_column_letter(index)].width = min(width, 70)
    with path.open("xb") as stream:
        workbook.save(stream)
    workbook.close()
